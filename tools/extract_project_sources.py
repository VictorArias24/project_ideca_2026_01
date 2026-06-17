from pathlib import Path

from docx import Document
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]

DOCX_FILES = [
    "Arquitectura.docx",
    "SOTA_VLM_IDECA_Classificacion_TIpos_De_Uso_2026.docx",
    "PLAN_E_IDEAS.docx",
]

XLSX_FILES = [
    "Victor Compendio de Usos para el Modelo.xlsx",
    "Victor Compendio de 10 Usos para el Modelo.xlsx",
]


def clean(value):
    if value is None:
        return ""
    return str(value).replace("\n", " / ").strip()


def print_docx_tables(path):
    document = Document(ROOT / path)
    print(f"\n### {path}")
    print(f"paragraphs={len(document.paragraphs)} tables={len(document.tables)}")
    for table_index, table in enumerate(document.tables, 1):
        print(f"\nTABLE {table_index}: rows={len(table.rows)} cols={len(table.columns)}")
        for row in table.rows[:25]:
            print(" | ".join(clean(cell.text)[:300] for cell in row.cells))
        if len(table.rows) > 25:
            print("...")


def print_xlsx(path):
    workbook = load_workbook(ROOT / path, data_only=True)
    print(f"\n### {path}")
    print(f"sheets={workbook.sheetnames}")
    for sheet in workbook.worksheets:
        print(f"\nSHEET {sheet.title}: rows={sheet.max_row} cols={sheet.max_column}")
        for row in sheet.iter_rows(
            min_row=1,
            max_row=min(sheet.max_row, 40),
            values_only=True,
        ):
            print(" | ".join(clean(value)[:300] for value in row))
        if sheet.max_row > 40:
            print("...")


for docx_file in DOCX_FILES:
    print_docx_tables(docx_file)

for xlsx_file in XLSX_FILES:
    print_xlsx(xlsx_file)
